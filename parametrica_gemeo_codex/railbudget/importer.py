"""Importação reprodutível dos dois arquivos, sem modificar as origens."""
from pathlib import Path
import hashlib
import json
import sqlite3
from openpyxl import load_workbook

FILES = {
    'vp': ('VPCODEX2026', 'EAP_Orcamento_Via_Permanente_2026.xlsx'),
    'p2': ('CODEX_PARAMETRICO_GERAL', 'EAP_Orcamento_Via_Permanente_2026_Parte2.xlsx'),
}

def extract(parent, output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    catalog, snapshots, inventory = {}, {}, []
    for origin, (folder, filename) in FILES.items():
        path = Path(parent) / folder / filename
        wf = load_workbook(path, read_only=True, data_only=False)
        wv = load_workbook(path, read_only=True, data_only=True)
        snapshots[origin] = {}
        for sf in wf:
            formulas = [list(r) for r in sf.values]
            values = [list(r) for r in wv[sf.title].values]
            snapshots[origin][sf.title] = {'formulas': formulas, 'values': values}
            inventory.append({'origin': origin, 'file': str(path), 'sheet': sf.title,
                              'rows': len(values), 'columns': max(map(len, values), default=0),
                              'formula_count': sum(isinstance(c, str) and c.startswith('=') for r in formulas for c in r),
                              'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
            if sf.title in ('SIEC_SERVIÇOS', 'SIEC_INSUMOS', 'Base Parte 1', 'Mercado e pendências'):
                for n, row in enumerate(values[4:], 5):
                    if not row or not row[0]:
                        continue
                    code = str(row[0])
                    if not code.startswith(('SIEC-', 'P1-', 'CARELLI-')):
                        continue
                    row = row + [None] * 10
                    source = 'SIEC' if code.startswith('SIEC-') else ('Provisão' if code.startswith('P1-') else 'Mercado')
                    kind = 'Insumos' if sf.title == 'SIEC_INSUMOS' else 'Serviços'
                    date = str(row[4]) if source == 'SIEC' else ('setembro/2026 declarado' if source == 'Provisão' else '13/09/2026')
                    key = f'{source}|{kind}|{date}|{code}'
                    cost = float(row[3]) if isinstance(row[3], (int, float)) else None
                    record = {'key': key, 'code': code, 'description': str(row[1] or ''),
                              'unit': str(row[2] or ''), 'price': cost, 'source': source,
                              'kind': kind, 'date': date, 'status': 'Preço disponível' if cost is not None else str(row[3] or 'Sem preço'),
                              'reference': str(row[5] or ''), 'provenance': []}
                    if key in catalog:
                        previous = catalog[key]
                        if any(previous[k] != record[k] for k in ('description', 'unit', 'price')):
                            raise ValueError(f'Conflito na base: {key}; não foi escolhido preço silenciosamente.')
                    else:
                        catalog[key] = record
                    catalog[key]['provenance'].append({'file': str(path), 'sheet': sf.title, 'row': n,
                                                       'original_file': row[5] if source == 'SIEC' else None,
                                                       'original_row': row[6] if source == 'SIEC' else None})
        wf.close()
        wv.close()
    db = output / 'catalog.sqlite'
    with sqlite3.connect(db) as con:
        con.execute('DROP TABLE IF EXISTS prices')
        con.execute('CREATE TABLE prices (key TEXT PRIMARY KEY, code TEXT, description TEXT, unit TEXT, price REAL, source TEXT, kind TEXT, date TEXT, status TEXT, reference TEXT, provenance TEXT)')
        con.executemany('INSERT INTO prices VALUES (?,?,?,?,?,?,?,?,?,?,?)',
                        [tuple(r[k] for k in ('key','code','description','unit','price','source','kind','date','status','reference')) + (json.dumps(r['provenance'],ensure_ascii=False),) for r in catalog.values()])
        con.execute('CREATE INDEX prices_code ON prices(code, source, date)')
    for name, data in [('snapshots', snapshots), ('inventory', inventory)]:
        (output / f'{name}.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return catalog, snapshots
