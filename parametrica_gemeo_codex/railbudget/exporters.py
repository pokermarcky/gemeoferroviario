"""Exportadores em memória. PDFs equivalentes via ReportLab, sem Office instalado."""
from io import BytesIO
import ast
import json
from pathlib import Path
import reportlab
from datetime import datetime
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from html import escape
from railbudget.localization import nome_variavel, formula_legivel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.pagebreak import Break
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, KeepTogether
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

NAVY='19364B'; TEAL='147D83'
PALETTE={'SIEC':'E2F3EB','SINAPI':'E5EDF9','SICRO':'EEE7F7','Mercado':'DDEEF6','Provisão':'FFF0CD'}

def font_pdf():
    pasta=Path(reportlab.__file__).parent/'fonts'
    if 'Vera' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Vera',str(pasta/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('VeraBd',str(pasta/'VeraBd.ttf')))
    return 'Vera','VeraBd'

def br(v, digits=2):
    return f'{v:,.{digits}f}'.replace(',','X').replace('.',',').replace('X','.')

def currency(v):return 'R$ '+br(v)

def caption(r):
    p=r['scenario']
    return f"{r['model_label']} | {p['configuration']} | {'dupla' if p['lines']==2 else 'simples'} | {br(p['km'],3)} km de corredor"

def excel_expr(expr, refs):
    def convert(n):
        if isinstance(n,ast.Constant):return str(n.value)
        if isinstance(n,ast.Name):return "'Premissas'!$B$"+str(refs[n.id])
        if isinstance(n,ast.BinOp):return '('+convert(n.left)+{ast.Add:'+',ast.Sub:'-',ast.Mult:'*',ast.Div:'/',ast.Pow:'^'}[type(n.op)]+convert(n.right)+')'
        if isinstance(n,ast.UnaryOp):return ('-' if isinstance(n.op,ast.USub) else '+')+convert(n.operand)
        if isinstance(n,ast.Call):
            name=n.func.id; args=[convert(a) for a in n.args]
            if name in ('ceil','floor'):return ('ROUNDUP' if name=='ceil' else 'ROUNDDOWN')+'('+args[0]+',0)'
            return name.upper()+'('+','.join(args)+')'
        raise ValueError('Expressão Excel não suportada')
    return '='+convert(ast.parse(expr.replace('^','**'),mode='eval').body)

def cache_formulas(data, caches):
    """Escreve resultados já validados em <v>; Excel recalcula normalmente ao editar."""
    out=BytesIO();ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(BytesIO(data)) as src,ZipFile(out,'w',ZIP_DEFLATED) as dest:
        for name in src.namelist():
            raw=src.read(name)
            if name in caches:
                tree=ET.fromstring(raw)
                for cell in tree.findall('.//s:c',ns):
                    address=cell.attrib['r']
                    if address not in caches[name]:continue
                    v=cell.find('s:v',ns)
                    if v is None:v=ET.SubElement(cell,'{'+ns['s']+'}v')
                    value=caches[name][address]
                    if isinstance(value,str):cell.set('t','str')
                    else:cell.attrib.pop('t',None)
                    v.text=str(value)
                raw=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            dest.writestr(name,raw)
    return out.getvalue()

def make_excel(r,catalog):
    w=Workbook();summary=w.active;summary.title='Resumo';eap=w.create_sheet('EAP');prem=w.create_sheet('Premissas')
    caches={f'xl/worksheets/sheet{i}.xml':{} for i in (1,2,3)}
    summary.append(['Parametric Rails | Orçamento']);summary.append([caption(r)])
    summary.append(['Grupo','Custo direto (R$)'])
    for g,v in r['groups'].items():summary.append([g,f'=SUMIF(EAP!$B$5:$B${4+len(r["items"])},A{summary.max_row+1},EAP!$I$5:$I${4+len(r["items"])})']);caches['xl/worksheets/sheet1.xml'][f'B{summary.max_row}']=v
    first_group_row=4;last_group_row=3+len(r['groups']);direct_row=last_group_row+1;bdi_row=direct_row+1
    bdi_amount_row=bdi_row+1;total_row=bdi_row+2;per_km_row=bdi_row+3;per_line_row=bdi_row+4
    summary.append(['Direto',f'=SUM(B{first_group_row}:B{last_group_row})']);caches['xl/worksheets/sheet1.xml'][f'B{direct_row}']=r['direct']
    summary.append(['BDI',r['scenario']['bdi']]);summary[f'B{bdi_row}'].number_format='0.00%'
    for label,formula,val in [('Adicional BDI',f'=ROUND(B{direct_row}*B{bdi_row},2)',r['bdi_amount']),('Total',f'=B{direct_row}+B{bdi_amount_row}',r['total']),('R$/km corredor',f'=B{total_row}/{r["scenario"]["km"]}',r['per_km']),('R$/km linha',f'=B{total_row}/{r["scenario"]["km"]*r["scenario"]["lines"]}',r['per_line_km'])]:
        summary.append([label,formula]);caches['xl/worksheets/sheet1.xml'][f'B{summary.max_row}']=val
    summary.append(['Prazo (meses)',r['duration']]);summary.append(['Drenagem',r['scenario']['drainage']]);summary.append(['Vedação',r['scenario']['fence']])
    for note in r['warnings']:summary.append([note])
    chart=BarChart();chart.title='Composição do custo direto';chart.add_data(Reference(summary,min_col=2,min_row=3,max_row=last_group_row),titles_from_data=True);chart.set_categories(Reference(summary,min_col=1,min_row=first_group_row,max_row=last_group_row))
    chart.height=11;chart.width=23;chart.legend=None;chart.title=None
    summary['D3']='Composição do custo direto (R$)'
    chart.x_axis.tickLblPos='nextTo';chart.y_axis.tickLblPos='nextTo';chart.x_axis.delete=False;chart.y_axis.delete=False
    summary.add_chart(chart,'D4')
    prem.append(['Parâmetro / variável','Valor / fórmula']);refs={name:i+2 for i,name in enumerate(r['context'])}
    derived=dict(r['derived'])
    summary[f'B{bdi_row}']="='Premissas'!B"+str(refs['bdi']);caches['xl/worksheets/sheet1.xml'][f'B{bdi_row}']=r['scenario']['bdi']
    summary[f'B{per_km_row}']=f"=B{total_row}/'Premissas'!B"+str(refs['km'])
    summary[f'B{per_line_row}']=f"=B{total_row}/('Premissas'!B"+str(refs['km'])+"*'Premissas'!B"+str(refs['lines'])+")"
    for name,val in r['context'].items():
        prem.append([nome_variavel(name),excel_expr(derived[name],refs) if name in derived else val]);caches['xl/worksheets/sheet3.xml'][f'B{prem.max_row}']=int(val) if isinstance(val,bool) else val
    sheets={};lookup={}
    used={x['price_key'] for x in r['items']}
    for k in sorted(used):
        p=catalog[k];name=p['source']+('_INSUMOS' if p['kind']=='Insumos' else '_SERVICOS') if p['source']=='SIEC' else p['source']
        if name not in sheets:
            sheets[name]=w.create_sheet(name);sheets[name].append(['Chave exata','Código','Descrição','Unidade','Preço','Data-base','Proveniência'])
        sh=sheets[name];sh.append([k,p['code'],p['description'],p['unit'],p['price'],p['date'],p['provenance']]);lookup[k]=name
    eap.append(['EAP orçada • '+caption(r)]);eap.append(['Quantidades vinculadas às premissas; preços por PROCV exato. Bases abaixo contêm referências usadas neste cenário.'])
    eap.append(['']);eap.append(['EAP','Grupo','Código','Fonte','Descrição / aplicação','Unidade','Quantidade','Custo unitário','Custo total','Data-base','Chave do preço','Memória de quantidade','Origem / memória'])
    for x in r['items']:
        n=eap.max_row+1;name=lookup[x['price_key']];span=f"'{name}'!$A$2:$G${sheets[name].max_row}"
        eap.append([x['eap'],x['group'],x['code'],x['source'],x['label']+' — '+x['description'],
                    f'=VLOOKUP(K{n},{span},4,FALSE)',excel_expr(x['quantity_formula'],refs),f'=VLOOKUP(K{n},{span},5,FALSE)',f'=ROUND(G{n}*H{n},2)',x['date'],x['price_key'],formula_legivel(x['quantity_formula']),f"{x['origin']} / {x['sheet']} / linha {x['row']}. {x['note']}"])
        for c,val in [('F',x['unit']),('G',x['quantity']),('H',x['unit_cost']),('I',x['total'])]:caches['xl/worksheets/sheet2.xml'][f'{c}{n}']=val
    for src,color in PALETTE.items():eap.conditional_formatting.add(f'A5:M{eap.max_row}',FormulaRule(formula=[f'$D5="{src}"'],fill=PatternFill('solid',fgColor=color)))
    for sh in w:
        header=4 if sh==eap else (3 if sh==summary else 1)
        sh.freeze_panes='A'+str(header+1);sh.auto_filter.ref=f'A{header}:{sh.cell(sh.max_row,sh.max_column).coordinate}'
        for cell in sh[header]:cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(name='Arial',bold=True,color='FFFFFF')
        for row in sh.iter_rows(min_row=header+1):
            for c in row:
                c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,(int,float)) or (c.data_type=='f'):c.number_format='#,##0.00'
        sh.sheet_view.showGridLines=False
        sh.sheet_properties.pageSetUpPr.fitToPage=True;sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
        sh.print_title_rows=f'1:{header}'
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=22
    summary.column_dimensions['A'].width=68;summary.column_dimensions['B'].width=24;summary[f'B{bdi_row}'].number_format='0.00%'
    summary.print_area='A1:Q30';summary.page_setup.fitToHeight=1
    summary.merge_cells('D3:Q3')
    eap.column_dimensions['E'].width=66;eap.column_dimensions['M'].width=78;eap.column_dimensions['K'].hidden=True;eap.column_dimensions['L'].width=45
    for row in range(5,eap.max_row+1):eap.row_dimensions[row].height=75;eap[f'G{row}'].number_format='#,##0.000000'
    prem.column_dimensions['A'].width=34;prem.column_dimensions['B'].width=30
    for sh in sheets.values():sh.column_dimensions['C'].width=75;sh.column_dimensions['G'].width=100
    add_group_sheets(w,r,refs,caches,{'direct':direct_row,'bdi':bdi_row,'total':total_row})
    data=BytesIO();w.save(data)
    return cache_formulas(data.getvalue(),caches)


def add_group_sheets(workbook,result,refs,caches,summary_rows):
    """Visões por grupo vinculadas à EAP consolidada, sem duplicar cálculos."""
    edge=Side(style='thin',color='CBD5E1')
    money_format='"R$" #,##0.00'
    for group,direct in result['groups'].items():
        sheet=workbook.create_sheet(group[:31])
        cache={};caches[f'xl/worksheets/sheet{len(workbook.worksheets)}.xml']=cache
        rows=[(n,x) for n,x in enumerate(result['items'],5) if x['group']==group]
        included=group in result.get('selected_groups',result['groups'])
        sheet['B2']=group;sheet.merge_cells('B2:J2')
        sheet['B3']=caption(result);sheet.merge_cells('B3:J3')
        sheet['B4']='Incluído no total selecionado' if included else 'Excluído do total selecionado'
        sheet.merge_cells('B4:J4')
        sheet['B5']='Valores vinculados à EAP consolidada. A seleção de escopo é feita na aplicação.'
        sheet.merge_cells('B5:J5')
        measures=[('Subtotal direto',f'=SUM(I13:I{12+len(rows)})' if rows else '=0',direct),
            ('Direto por km de corredor',f"=C6/'Premissas'!B{refs['km']}",direct/result['scenario']['km']),
            ('Participação no total direto',f'=IF(Resumo!B{summary_rows["direct"]}=0,0,C6/Resumo!B{summary_rows["direct"]})',direct/result['direct'] if result['direct'] else 0),
            ('BDI',f'=Resumo!B{summary_rows["bdi"]}',result['scenario']['bdi'])]
        for n,(label,formula,value) in enumerate(measures,6):
            sheet.cell(n,2,label);sheet.cell(n,3,formula);cache[f'C{n}']=value
        headers=['EAP','Código','Fonte','Descrição do serviço','Unidade','Quantidade','Custo unitário','Custo total','Data-base']
        for col,label in enumerate(headers,2):sheet.cell(12,col,label)
        source_cols=['A','C','D','E','F','G','H','I','J']
        for n,(source_row,item) in enumerate(rows,13):
            values=[item['eap'],item['code'],item['source'],item['label']+' — '+item['description'],item['unit'],item['quantity'],item['unit_cost'],item['total'],item['date']]
            for col,source_col,value in zip(range(2,11),source_cols,values):
                sheet.cell(n,col,f"='EAP'!{source_col}{source_row}");cache[f'{get_column_letter(col)}{n}']=value
            sheet.row_dimensions[n].height=100
        if not rows:sheet['B13']='Sem serviços incluídos neste grupo.';sheet.merge_cells('B13:J13')
        for row in sheet.iter_rows(min_row=2,min_col=2,max_row=max(13,12+len(rows)),max_col=10):
            for cell in row:
                cell.font=Font(name='Aptos',size=12,color=NAVY)
                cell.alignment=Alignment(horizontal='left' if cell.column==5 else 'center',vertical='center',wrap_text=True)
                if cell.row>=12 or (6<=cell.row<=9 and cell.column<=3):cell.border=Border(left=edge,right=edge,top=edge,bottom=edge)
        for cell in sheet[12][1:10]:cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(name='Aptos',size=12,bold=True,color='FFFFFF')
        sheet['B2'].font=Font(name='Aptos',size=16,bold=True,color=NAVY)
        for n in (6,7):sheet[f'C{n}'].number_format=money_format
        for n in (8,9):sheet[f'C{n}'].number_format='0.00%'
        for n in range(13,13+len(rows)):
            sheet[f'G{n}'].number_format='#,##0.000000'
            for col in ('H','I'):sheet[f'{col}{n}'].number_format=money_format
        for col,width in {'A':3,'B':30,'C':27,'D':14,'E':76,'F':14,'G':22,'H':24,'I':24,'J':18}.items():sheet.column_dimensions[col].width=width
        for n in range(2,13):sheet.row_dimensions[n].height=32 if n==12 else 28
        sheet.freeze_panes='G13';sheet.sheet_view.showGridLines=False
        sheet.auto_filter.ref=f'B12:J{max(12,12+len(rows))}'
        sheet.print_title_rows='2:12';sheet.print_area=f'B2:J{max(13,12+len(rows))}'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0

def make_word(r):
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
    style=d.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for s in ['Title','Heading 1','Heading 2']:d.styles[s].font.name='Arial';d.styles[s].font.color.rgb=RGBColor.from_string(NAVY)
    d.add_heading('Orçamento paramétrico ferroviário',0);d.add_paragraph(caption(r));d.add_paragraph('Regra '+r['rule_version']+' • '+datetime.now().strftime('%d/%m/%Y'))
    d.add_heading('Resultado do cenário',1)
    def table(headers,rows,widths):
        t=d.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for row in t.rows:
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(9)
        for cell in t.rows[0].cells:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY);cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
        return t
    table(['Grupo','Direto'],[[g,currency(v)] for g,v in r['groups'].items()]+[['Total com BDI',currency(r['total'])],['R$/km corredor',currency(r['per_km'])],['R$/km linha',currency(r['per_line_km'])]],[10,7])
    p=r['scenario'];d.add_paragraph(f"Drenagem {p['drainage']}; vedação {p['fence']}; {p['amvs']} AMV(s) no corredor; banco de dutos {'sim' if p['ducts'] else 'não'}; topografia {'sim' if p['topography'] else 'não'}; rede aérea {'sim' if p['overhead'] else 'não'}; sinalização {'sim' if p['signaling'] else 'não'} ({p['detection'].lower()}); material rodante {'sim' if p['rolling_stock'] else 'não'} ({p['trainsets']} composição(ões)); prazo {r['duration']} meses; BDI {br(p['bdi']*100,6)}%.")
    d.add_heading('Fontes e limites',1)
    d.add_paragraph('; '.join(f'{k}: {v} linhas' for k,v in r['counts'].items()))
    for note in r['warnings']:d.add_paragraph(note)
    d.add_paragraph('Preços SIEC de junho/2026. Material rodante baseado no contrato CPTM 8186142011, data-base abril/2016, sem reajuste. Código, data-base, aba e linha são preservados na memória detalhada e no Excel; caminhos locais e endereços eletrônicos não são publicados.')
    d.add_page_break();d.add_heading('Memória de quantidades e preços',1)
    for g in r['groups']:
        rows=[x for x in r['items'] if x['group']==g]
        if not rows:continue
        d.add_heading(g,2)
        for x in rows:
            d.add_paragraph(f"{x['eap']} | {x['label']} | {x['code']} ({x['source']}, {x['date']})",style='Heading 3')
            detail=d.add_paragraph(f"{br(x['quantity'],6)} {x['unit']} × {currency(x['unit_cost'])} = {currency(x['total'])}. Memória de cálculo: {formula_legivel(x['quantity_formula'])}. Origem: {x['origin']} / {x['sheet']} / linha {x['row']}.")
            detail.paragraph_format.keep_with_next=True
            d.add_paragraph(x['note'])
    out=BytesIO();d.save(out);return out.getvalue()

def make_pdf(r,detailed=False):
    out=BytesIO();page=landscape(A3) if detailed else A4
    regular,bold=font_pdf()
    styles=getSampleStyleSheet()
    for title in ('Title','Heading1','Heading2','Heading3'):styles[title].fontName=bold
    styles.add(ParagraphStyle(name='CellR',fontName=regular,fontSize=8 if detailed else 9,leading=12))
    styles.add(ParagraphStyle(name='HeadR',fontName=bold,fontSize=8 if detailed else 9,leading=12,textColor=colors.white))
    def p(text,style='CellR'):return Paragraph(escape(str(text)),styles[style])
    def t(headers,rows,widths):
        obj=Table([[p(h,'HeadR') for h in headers]]+[[p(v) for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        obj.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+NAVY)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F0F5F7')]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));return obj
    flow=[p('EAP e orçamento' if detailed else 'Relatório técnico do orçamento','Title'),Spacer(1,10),p(caption(r)),Spacer(1,12)]
    flow.append(t(['Grupo','Custo direto'],[[g,currency(v)] for g,v in r['groups'].items()]+[['Total com BDI',currency(r['total'])],['R$/km corredor',currency(r['per_km'])],['R$/km linha',currency(r['per_line_km'])]],[340,180] if not detailed else [650,470]))
    sc=r['scenario'];flow += [Spacer(1,14),p(f"Drenagem: {sc['drainage']} | Vedação: {sc['fence']} | AMVs: {sc['amvs']} | Dutos: {'sim' if sc['ducts'] else 'não'} | Topografia: {'sim' if sc['topography'] else 'não'} | Rede aérea: {'sim' if sc['overhead'] else 'não'} | Sinalização: {'sim' if sc['signaling'] else 'não'} ({sc['detection'].lower()}) | Material rodante: {sc['trainsets'] if sc['rolling_stock'] else 0} composição(ões) | Prazo: {r['duration']} meses | BDI: {br(sc['bdi']*100,6)}%"),Spacer(1,10)]
    for note in r['warnings']:flow.extend([p(note),Spacer(1,7)])
    flow.append(p('Fontes por linha: '+'; '.join(f'{k}: {v}' for k,v in r['counts'].items())))
    flow.append(p('Origem: VPCODEX2026 + CODEX_PARAMETRICO_GERAL. SIEC junho/2026. Regra '+r['rule_version']))
    if detailed:
        flow.append(PageBreak());flow.append(p('EAP detalhada','Heading1'))
        rows=[[x['eap'],x['group'],x['code'],x['source'],x['label']+' — '+x['description'],x['unit'],br(x['quantity'],6),currency(x['unit_cost']),currency(x['total'])] for x in r['items']]
        flow.append(t(['EAP','Grupo','Código','Fonte','Descrição','Un.','Quantidade','Unitário','Total'],rows,[65,90,130,60,365,40,90,90,90]))
    else:
        flow.append(PageBreak());flow.append(p('Memória de quantidades e preços','Heading1'))
        for x in r['items']:
            flow.append(KeepTogether([p(x['group']+' / '+x['label'],'Heading3'),p(f"{x['code']} | {x['source']} | {x['date']}"),p(f"{br(x['quantity'],6)} {x['unit']} × {currency(x['unit_cost'])} = {currency(x['total'])}"),p('Memória de cálculo: '+formula_legivel(x['quantity_formula'])),p(f"Origem: {x['origin']} / {x['sheet']} / linha {x['row']}. "+x['note']),Spacer(1,8)]))
    def footer(c,d):c.setFont(regular,8);c.drawString(32,18,'Gêmeo ferroviário • Orçamento paramétrico');c.drawRightString(page[0]-32,18,str(d.page))
    SimpleDocTemplate(out,pagesize=page,leftMargin=32,rightMargin=32,topMargin=32,bottomMargin=34).build(flow,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()

def export_all(r,catalog):
    return {'orcamento.xlsx':make_excel(r,catalog),'relatorio.docx':make_word(r),
            'orcamento.pdf':make_pdf(r,True),'relatorio.pdf':make_pdf(r,False)}
