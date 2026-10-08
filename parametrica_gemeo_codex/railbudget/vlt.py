"""Modelo paramétrico inicial do Ramal VLT Aeroporto - Castelão.

Os valores reproduzem integralmente os totais por grupo do Anexo A do edital,
data-base fevereiro/2025. As extensões de corredor são inferidas da montagem de
1.800 m de via em superfície e 3.400 m de via LVT elevada, ambas em via dupla.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from railbudget.expressions import money


BASE_SURFACE_KM=0.90
BASE_ELEVATED_KM=1.70
BASE_CORRIDOR_KM=2.60
BASE_LINE_KM=5.20
BASE_TOTAL=191_008_605.92
REFERENCE_PERIOD="fevereiro/2025"

# (EAP, descrição, critério paramétrico, total do edital com BDI)
BASE_GROUPS=(
    ("1","Serviços preliminares","Fixo",9_435_849.21),
    ("2","Administração local","Fixo",5_742_873.68),
    ("3","Movimento de terra","Superfície",2_820_491.10),
    ("4","Estruturas de concreto armado, protendido e metálica","Elevado",96_111_086.49),
    ("5","Urbanização","Superfície",6_565_436.71),
    ("6","Acabamentos","Comum",4_258_024.50),
    ("7","Comunicação visual","Comum",94_379.34),
    ("8","Serviços complementares","Comum",7_276.80),
    ("9","Instalações hidrossanitárias","Comum",2_559_328.33),
    ("10","Instalações elétricas","Comum",753_661.93),
    ("11","Equipamentos","Comum",27_349.02),
    ("12.S","Via permanente em superfície, passagens de nível e segurança","Superfície",11_164_773.46),
    ("12.E","Via permanente LVT em elevado","Elevado",14_988_189.01),
    ("13","Sistemas fixos e de controle","Comum",36_479_886.34),
)


def calculate_vlt(surface_km=BASE_SURFACE_KM,elevated_km=BASE_ELEVATED_KM):
    surface_km=float(surface_km);elevated_km=float(elevated_km)
    if surface_km<0 or elevated_km<0 or surface_km+elevated_km<=0:
        raise ValueError("Informe uma extensão positiva em superfície ou elevado.")
    corridor_km=surface_km+elevated_km
    factors={"Fixo":1.0,"Superfície":surface_km/BASE_SURFACE_KM,
        "Elevado":elevated_km/BASE_ELEVATED_KM,"Comum":corridor_km/BASE_CORRIDOR_KM}
    items=[];groups={};segments={"Custos fixos":0.0,"Trecho em superfície":0.0,
        "Trecho elevado":0.0,"Sistemas e instalações comuns":0.0}
    segment_for={"Fixo":"Custos fixos","Superfície":"Trecho em superfície",
        "Elevado":"Trecho elevado","Comum":"Sistemas e instalações comuns"}
    for sequence,(eap,label,basis,baseline) in enumerate(BASE_GROUPS,1):
        factor=factors[basis];total=money(baseline*factor);group=f"{eap} {label}"
        groups[group]=total;segments[segment_for[basis]]=money(segments[segment_for[basis]]+total)
        items.append({"id":f"VLT-{eap}","eap":eap,"group":group,"label":label,
            "code":f"GRUPO {eap}","source":"Edital VLT","description":label,
            "unit":"VB","quantity":factor,"unit_cost":baseline,"total":total,
            "date":REFERENCE_PERIOD,"basis":basis,"baseline_cost":baseline,
            "note":"Total com BDI do Anexo A; fator aplicado conforme o tipo de implantação."})
    total=money(sum(groups.values()));line_km=corridor_km*2
    return {"scenario":{"configuration":"Híbrido - superfície + elevado","km":corridor_km,
            "surface_km":surface_km,"elevated_km":elevated_km,"lines":2,"line_km":line_km,
            "bdi":0.0,"source_bdi_included":True},
        "model_label":"VLT Aeroporto - Castelão | Edital","reference_period":REFERENCE_PERIOD,
        "items":items,"groups":groups,"segments":segments,"direct":total,"bdi_amount":0.0,
        "total":total,"per_km":money(total/corridor_km),"per_line_km":money(total/line_km),
        "source_total":BASE_TOTAL,"rule_version":"VLT-EDITAL-FEV2025-1",
        "counts":{"Edital VLT":len(items)},
        "warnings":[
            "Valores com BDI já incorporado conforme o edital: materiais 16,80%, serviços 24,23% e betuminosos 15,00%.",
            "Extensões do corredor inferidas para via dupla: 1.800 m de montagem em superfície equivalem a 0,90 km de corredor; 3.400 m de LVT equivalem a 1,70 km de corredor elevado.",
            "Custos fixos permanecem constantes; superfície e elevado variam por suas extensões; sistemas e instalações comuns variam pela extensão total.",
            "Material rodante não aparece como grupo autônomo na planilha e não foi acrescentado por fonte externa.",
        ]}


def make_vlt_excel(result):
    workbook=Workbook();summary=workbook.active;summary.title="Resumo"
    detail=workbook.create_sheet("Grupos do edital");premises=workbook.create_sheet("Premissas")
    summary.append(["Parametric Rails | Orçamento VLT"])
    summary.append(["Projeto","VLT Aeroporto - Castelão"])
    summary.append(["Data-base",result["reference_period"]])
    summary.append(["Implantação",result["scenario"]["configuration"]])
    summary.append(["Extensão em superfície (km)",result["scenario"]["surface_km"]])
    summary.append(["Extensão elevada (km)",result["scenario"]["elevated_km"]])
    summary.append(["Extensão total do corredor (km)",result["scenario"]["km"]])
    summary.append(["Extensão total de linha - via dupla (km)",result["scenario"]["line_km"]])
    summary.append(["Custo total",result["total"]])
    summary.append(["Custo por km de corredor",result["per_km"]])
    summary.append(["Custo por km de linha",result["per_line_km"]])
    summary.append([]);summary.append(["Componente","Valor"])
    for label,value in result["segments"].items():summary.append([label,value])

    detail.append(["EAP","Grupo do edital","Aplicação","Custo-base com BDI","Fator paramétrico","Custo no cenário","Data-base"])
    for item in result["items"]:
        detail.append([item["eap"],item["label"],item["basis"],item["baseline_cost"],
            item["quantity"],item["total"],item["date"]])

    premises.append(["Premissas do modelo VLT"])
    premises.append(["Premissa","Valor","Unidade/critério"])
    premises.append(["Fonte","Anexo A - Planilha de Preços Básicos","Edital"])
    premises.append(["Data-base",REFERENCE_PERIOD,""])
    premises.append(["Traçado-base em superfície",BASE_SURFACE_KM,"km de corredor em via dupla"])
    premises.append(["Traçado-base elevado",BASE_ELEVATED_KM,"km de corredor em via dupla"])
    premises.append(["Via em superfície da planilha",1.80,"km de linha"])
    premises.append(["Via LVT elevada da planilha",3.40,"km de linha"])
    premises.append(["BDI de materiais",0.168,"já incorporado"])
    premises.append(["BDI de serviços",0.2423,"já incorporado"])
    premises.append(["BDI de betuminosos",0.15,"já incorporado"])
    premises.append(["Total de controle do edital",BASE_TOTAL,"R$"])
    for warning in result["warnings"]:premises.append(["Observação",warning,""])

    navy="19364B";teal="147D83"
    for sheet,header in ((summary,13),(detail,1),(premises,2)):
        sheet.freeze_panes=None;sheet.sheet_view.showGridLines=False
        sheet.auto_filter.ref=f"A{header}:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}"
        for cell in sheet[header]:
            cell.fill=PatternFill("solid",fgColor=navy);cell.font=Font(name="Aptos",bold=True,color="FFFFFF")
        for row in sheet.iter_rows():
            for cell in row:
                cell.font=Font(name="Aptos",size=10,bold=cell.font.bold,color=cell.font.color)
                cell.alignment=Alignment(vertical="top",wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation="landscape"
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    summary["A1"].font=Font(name="Aptos Display",size=16,bold=True,color=navy)
    summary.column_dimensions["A"].width=48;summary.column_dimensions["B"].width=32
    for row in (9,10,11,14,15,16,17):summary[f"B{row}"].number_format='[$R$-pt-BR] #,##0.00'
    detail.column_dimensions["B"].width=68
    for column in ("D","F"):
        for cell in detail[column][1:]:cell.number_format='[$R$-pt-BR] #,##0.00'
    detail.column_dimensions["C"].width=22;detail.column_dimensions["G"].width=20
    premises.column_dimensions["A"].width=34;premises.column_dimensions["B"].width=78;premises.column_dimensions["C"].width=28
    for cell in premises[1]:cell.font=Font(name="Aptos Display",size=16,bold=True,color=teal)
    output=BytesIO();workbook.save(output);return output.getvalue()
