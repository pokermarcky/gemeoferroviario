"""Comparação lado a lado de dois orçamentos já calculados.

Recebe resultados prontos (após o recorte de grupos e BDI) e só lê seus
valores: não recalcula nem altera nenhum orçamento.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .expressions import money

INDICATORS = (
    ('direct', 'Custo direto'),
    ('bdi_amount', 'Acréscimo de BDI'),
    ('total', 'Total'),
    ('per_km', 'Por km de corredor'),
    ('per_line_km', 'Por km de linha'),
    ('per_track_km', 'Por km de via'),
    ('fleet_total', 'Frota (material rodante)'),
    ('per_km_without_fleet', 'Infraestrutura por km'),
)


def _row(item, kind, a, b):
    a = float(a or 0.0)
    b = float(b or 0.0)
    difference = money(b - a)
    return {'item': item, 'kind': kind, 'a': money(a), 'b': money(b), 'difference': difference,
            'percent': (100 * difference / a) if a else None}


def compare_results(a, b, label=lambda group: group):
    """Linhas da comparação: grupos incluídos em A ou B, depois os indicadores.

    ``percent`` é a variação de B sobre A; fica ``None`` quando A é zero.
    """
    selected_a = set(a.get('selected_groups', a['groups']))
    selected_b = set(b.get('selected_groups', b['groups']))
    groups = [g for g in a['groups'] if g in selected_a or g in selected_b]
    groups += [g for g in b['groups'] if g not in a['groups'] and g in selected_b]
    rows = [_row(label(g), 'group', a['groups'].get(g, 0.0) if g in selected_a else 0.0,
                 b['groups'].get(g, 0.0) if g in selected_b else 0.0) for g in groups]
    for key, name in INDICATORS:
        if a.get(key) is None and b.get(key) is None:
            continue
        if key == 'fleet_total' and not (a.get(key) or b.get(key)):
            continue
        if key == 'per_km_without_fleet' and not (a.get('fleet_total') or b.get('fleet_total')):
            continue
        rows.append(_row(name, 'indicator', a.get(key), b.get(key)))
    return rows


def comparison_excel(rows, caption_a, caption_b):
    """Planilha simples com a comparação, em valores fixos."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Comparação'
    sheet.append(['Parametric Rails | Comparação de cenários'])
    sheet.append(['Cenário A', caption_a])
    sheet.append(['Cenário B', caption_b])
    sheet.append([])
    header = ['Item', 'Cenário A (R$)', 'Cenário B (R$)', 'Diferença B − A (R$)', 'Variação (%)']
    sheet.append(header)
    fill = PatternFill('solid', fgColor='B8322A')
    for cell in sheet[5]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
    sheet['A1'].font = Font(bold=True, size=14)
    for row in rows:
        sheet.append([row['item'], row['a'], row['b'], row['difference'],
                      None if row['percent'] is None else row['percent'] / 100])
        current = sheet.max_row
        for column in 'BCD':
            sheet[f'{column}{current}'].number_format = '"R$" #,##0.00'
        sheet[f'E{current}'].number_format = '0.00%'
        if row['kind'] == 'indicator':
            sheet[f'A{current}'].font = Font(bold=True)
    for column, width in zip('ABCDE', (42, 22, 22, 24, 14)):
        sheet.column_dimensions[column].width = width
    out = BytesIO()
    book.save(out)
    return out.getvalue()
