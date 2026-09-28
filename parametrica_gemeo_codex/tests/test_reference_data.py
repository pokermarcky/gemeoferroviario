from railbudget.engine import load_model
from io import BytesIO

from openpyxl import Workbook, load_workbook

from railbudget.reference_data import parse_reference, apply_reference_bases, normalized_excel


def test_uploaded_base_replaces_linked_price_and_keeps_only_current_slot():
    _, catalog = load_model()
    target = next(item for item in catalog.values()
                  if item['source'] == 'SIEC' and item['kind'] == 'Serviços' and item['price'])
    csv = ('Código;Descrição;Unidade;Preço\n'
           f'{target["code"]};Preço atualizado;{target["unit"]};4.079.193,32\n').encode()
    base = parse_reference(csv, 'siec-08-2026.csv', 'SIEC', 'Serviços', '08/2026')
    updated, linked = apply_reference_bases(catalog, {('SIEC', 'Serviços'): base})
    assert linked[('SIEC', 'Serviços')] >= 1
    assert updated[target['key']]['price'] == 4079193.32
    assert updated[target['key']]['date'] == '08/2026'


def test_reference_requires_pricing_columns():
    try:
        parse_reference(b'codigo;texto\n1;sem preco\n', 'base.csv', 'SINAPI', 'Insumos', '08/2026')
    except ValueError as exc:
        assert 'Colunas obrigatórias' in str(exc)
    else:
        raise AssertionError('A tabela incompleta deveria ser rejeitada.')


def test_siurb_excel_with_cover_and_displaced_header_is_converted():
    book = Workbook()
    cover = book.active
    cover.title = 'CAPA'
    cover['A1'] = 'Tabela de custos SIURB'
    data = book.create_sheet('INSUMOS')
    data.append(['Prefeitura de São Paulo'])
    data.append(['Data-base: 07/2026'])
    data.append([])
    data.append(['Código do Insumo', 'Discriminação', 'Unid.', 'Preço Unitário'])
    data.append(['MAT-001', 'Brita graduada', 'm³', '4.079,32'])
    data.append(['MAT-002', 'Dormente', 'un', 187.5])
    stream = BytesIO()
    book.save(stream)

    base = parse_reference(stream.getvalue(), 'siurb-07-2026.xlsx', 'SIURB', 'Insumos', '07/2026')

    assert base.sheets == ['INSUMOS']
    assert base.count == 2
    assert base.records[0]['price'] == 4079.32
    assert base.records[1]['code'] == 'MAT-002'

    converted = load_workbook(BytesIO(normalized_excel(base)), data_only=True)
    sheet = converted['Base normalizada']
    assert [cell.value for cell in sheet[1]] == [
        'Fonte', 'Tipo', 'Código', 'Descrição', 'Unidade', 'Preço', 'Data-base']
    assert sheet['F2'].value == 4079.32


def test_excel_combines_valid_sheets_and_deduplicates_codes():
    book = Workbook()
    first = book.active
    first.title = 'Materiais A'
    first.append(['Código', 'Descrição', 'Unidade', 'Preço'])
    first.append(['001', 'Trilho', 'm', 100])
    second = book.create_sheet('Materiais B')
    second.append(['Código', 'Descrição', 'Unidade', 'Preço'])
    second.append(['002', 'Lastro', 'm³', 200])
    second.append(['001', 'Trilho atualizado', 'm', 150])
    stream = BytesIO()
    book.save(stream)

    base = parse_reference(stream.getvalue(), 'sicro.xlsx', 'SICRO', 'Insumos', '08/2026')

    assert base.sheets == ['Materiais A', 'Materiais B']
    assert base.count == 2
    assert {record['code']: record['price'] for record in base.records} == {'001': 150, '002': 200}
