from pathlib import Path

from streamlit.testing.v1 import AppTest


APP = Path(__file__).resolve().parents[1] / 'app.py'


def test_passageiro_calcula_siec_e_limpa_resultado_invalido():
    from railbudget.engine import Scenario, calculate, load_model

    app = AppTest.from_file(str(APP)).run(timeout=30)
    assert not app.exception
    app.button(key='main_calculate').click().run(timeout=30)
    assert app.session_state['results']['main']['total'] == calculate(Scenario(), *load_model())['total']
    app.number_input(key='main_km').set_value(.01)
    app.button(key='main_calculate').click().run(timeout=30)
    assert app.error
    assert 'main' not in app.session_state['results']


def test_navegacao_e_controles_solicitados():
    app = AppTest.from_file(str(APP)).run(timeout=30)
    assert not app.exception
    assert [tab.label for tab in app.tabs] == [
        'ORÇAMENTOS',
        'Ferrovia de passageiro', 'Resumo geral', 'Detalhamento por grupo',
        'Ferrovia de carga', 'Resumo geral', 'Detalhamento por grupo',
        'VLT - Veículo leve sobre Trilho', 'Shortline',
        'BASES DE REFERÊNCIA', 'Insumos', 'Serviços']
    assert any(h.value == 'Parametric Rails' for h in app.title)
    assert app.selectbox(key='main_profile').options == ['SIEC • lastro / AMV nº 14']
    assert app.selectbox(key='main_configuration').options == ['Superfície', 'Elevado', 'Subterrâneo']
    assert app.checkbox(key='main_grupo_5').label == 'Banco de dutos'
    assert not any(x.label == 'Prazo da obra (meses)' for x in app.number_input)
    assert [(x.key, x.proto.label) for x in app.get('download_button')] == [
        ('main_excel', 'Baixar orçamento em Excel'), ('cargo_excel', 'Baixar orçamento em Excel')]
    assert not any(x.key == 'main_prepare' for x in app.button)
    assert len(app.get('file_uploader')) == 6
    assert app.button(key='main_selecionar_todas').proto.type != 'primary'


def test_subterraneo_nao_reaproveita_preco_de_superficie():
    app = AppTest.from_file(str(APP)).run(timeout=30)
    app.selectbox(key='main_configuration').set_value('Subterrâneo').run(timeout=30)
    assert not app.exception
    assert app.button(key='main_calculate').proto.disabled
    assert not any(x.key == 'main_excel' for x in app.get('download_button'))


def test_selecao_de_grupos_e_bdi_no_resultado():
    from railbudget.expressions import money
    from railbudget.exporters import currency

    app = AppTest.from_file(str(APP)).run(timeout=30)
    direct = app.session_state['results']['main']['direct']
    app.number_input(key='main_bdi_personalizado').set_value(25.0).run(timeout=30)
    assert app.metric[0].value == currency(money(direct + money(direct * .25)))
    app.checkbox(key='main_aplicar_bdi').uncheck().run(timeout=30)
    assert app.metric[0].value == currency(direct)
    app.button(key='main_desmarcar_todas').click().run(timeout=30)
    assert app.metric[0].value == 'R$ 0,00'
    assert not any(x.key == 'main_excel' for x in app.get('download_button'))
    app.button(key='main_selecionar_todas').click().run(timeout=30)
    assert app.metric[0].value == currency(direct)
    assert len(app.get('download_button')) == 2


def test_banco_de_dutos_e_opcoes_do_grupo():
    app = AppTest.from_file(str(APP)).run(timeout=30)
    completo = app.metric[0].value
    app.checkbox(key='main_grupo_5').uncheck().run(timeout=30)
    assert not app.exception
    assert app.metric[0].value != completo
    app.checkbox(key='main_grupo_5').check().run(timeout=30)
    assert app.metric[0].value == completo
    app.selectbox(key='main_configuration').set_value('Elevado').run(timeout=30)
    assert app.checkbox(key='main_grupo_5').label == 'Canaletas e passa-fios'
    app.checkbox(key='main_grupo_3').uncheck().run(timeout=30)
    assert not any(x.key == 'main_fence' for x in app.selectbox)


def test_estacoes_precisam_preco_e_entram_no_excel():
    from io import BytesIO
    from openpyxl import load_workbook
    from railbudget.engine import select_groups
    from railbudget.exporters import make_excel
    from railbudget.engine import load_model
    app=AppTest.from_file(str(APP)).run(timeout=30)
    app.number_input(key='main_station_0').set_value(2).run(timeout=30)
    assert app.button(key='main_calculate').proto.disabled
    assert not any(x.key == 'main_excel' for x in app.get('download_button'))
    app.number_input(key='main_station_price_0').set_value(1200000.0).run(timeout=30)
    assert not any(x.key == 'main_excel' for x in app.get('download_button'))  # orçamento antigo oculto até recalcular
    app.button(key='main_calculate').click().run(timeout=30)
    assert not app.exception
    result=app.session_state['results']['main']
    assert result['groups']['10 Estações de passageiros']==2400000
    chosen=list(result['groups'])
    scoped=select_groups(result,chosen)
    book=load_workbook(BytesIO(make_excel(scoped,{**load_model()[1],**result['custom_prices']})),data_only=True)
    assert any(row[2]=='EST-01' and row[6]==2 and row[8]==2400000 for row in book['EAP'].iter_rows(min_row=5,values_only=True))


def test_carga_orca_so_infraestrutura_com_siec():
    app=AppTest.from_file(str(APP)).run(timeout=30)
    assert app.checkbox(key='cargo_grupo_5').value is False
    assert app.checkbox(key='cargo_grupo_4').value is False
    assert app.checkbox(key='cargo_grupo_6').value is False
    assert app.checkbox(key='cargo_grupo_7').value is False
    app.number_input(key='cargo_wagons').set_value(80).run(timeout=30)
    assert not app.exception
    cargo=app.session_state['results']['cargo']
    passenger=app.session_state['results']['main']
    assert cargo['total']<passenger['total']
    assert cargo['scenario']['rolling_stock'] is False
    assert '9 Material rodante' not in cargo['groups']
    assert all(x['source']=='SIEC' for x in cargo['items'])
    assert any('80 vagão' in warning for warning in cargo['warnings'])
    assert {x.key for x in app.get('download_button')}=={'main_excel','cargo_excel'}
