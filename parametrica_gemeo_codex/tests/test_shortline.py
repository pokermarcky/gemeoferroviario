from io import BytesIO

import pytest
from openpyxl import load_workbook

from railbudget.engine import Scenario, load_model, select_groups
from railbudget.exporters import make_excel
from railbudget.freight import calculate_freight
from railbudget.shortline import (SHORTLINE_GROUPS, SIDING_GROUP, ShortlineParams,
    calculate_shortline, default_selection, params_from_scenario)


@pytest.fixture(scope='module')
def model():
    return load_model()


def test_shortline_reaproveita_carga_siec_sem_preco_novo(model):
    rules,catalog=model
    result=calculate_shortline(ShortlineParams(),rules,catalog)
    assert list(result['groups'])==list(SHORTLINE_GROUPS)
    assert sum(i['total'] for i in result['items'])==pytest.approx(result['direct'],abs=.05)
    assert all(i['source']=='SIEC' for i in result['items'])
    assert all(catalog[i['price_key']]['price']==i['unit_cost'] for i in result['items'])
    assert not any(g.startswith(('6 ','7 ','9 ')) for g in result['groups'])
    rail=next(i for i in result['items'] if i['id']=='vp:Superfície:26')
    assert rail['code']=='SIEC-03.04.01.100.10'
    assert result['scenario']['amvs']==4 and result['scenario']['lines']==1


def test_linha_principal_bate_com_a_ferrovia_de_carga(model):
    rules,catalog=model
    params=ShortlineParams(km=5,sidings=0,yard_length=0,yard_amvs=1)
    result=calculate_shortline(params,rules,catalog)
    cargo=calculate_freight(Scenario(km=5,lines=1,drainage='Normal',amvs=1,ducts=False,
        overhead=False,detection='Contador de eixos',rolling_stock=False,trainsets=0),25,rules,catalog)
    assert result['groups'][SIDING_GROUP]==0
    for group in ('1 Via permanente','2 Topografia','3 Drenagem','4 Vedação','5 AMVs','8 Sinalização'):
        assert result['groups'][group]==cargo['groups'][group]


def test_desvios_e_patio_somam_via_sem_duplicar_canteiro(model):
    rules,catalog=model
    base=calculate_shortline(ShortlineParams(sidings=0,yard_length=0,yard_amvs=0),rules,catalog)
    more=calculate_shortline(ShortlineParams(sidings=2,siding_length=1000,yard_length=500,yard_amvs=0),rules,catalog)
    siding=[i for i in more['items'] if i['group']==SIDING_GROUP]
    assert siding and not any(i['label'] in ('Canteiro','Sondagens','Administração local','Mobilização') for i in siding)
    assert next(i for i in siding if i['label']=='Montagem')['quantity']==pytest.approx(2500)
    assert more['scenario']['amvs']==4 and more['scenario']['track_km']==pytest.approx(12.5)
    assert more['groups'][SIDING_GROUP]>0 and more['total']>base['total']
    assert more['groups']['1 Via permanente']<base['groups']['1 Via permanente']  # envelopes dos AMVs
    assert all(i['eap'].startswith('1.3.') for i in siding)


def test_escopo_pendente_fica_fora_do_total(model):
    rules,catalog=model
    plain=calculate_shortline(ShortlineParams(),rules,catalog)
    scoped=calculate_shortline(ShortlineParams(level_crossings=7,locomotives=3,wagons=40),rules,catalog)
    assert scoped['total']==plain['total']
    pending={row['item']:row['quantity'] for row in scoped['pending_scope']}
    assert pending['Passagens em nível']==7 and pending['Locomotivas']==3 and pending['Vagões']==40


def test_selecao_inicial_bdi_e_validacao(model):
    rules,catalog=model
    result=calculate_shortline(ShortlineParams(),rules,catalog)
    chosen=[g for g in result['groups'] if default_selection(g)]
    assert '8 Sinalização' not in chosen
    cut=select_groups(result,chosen,bdi_rate=.2)
    assert cut['total']==pytest.approx(cut['direct']*1.2,abs=.05)
    with pytest.raises(ValueError):calculate_shortline(ShortlineParams(km=0),rules,catalog)
    with pytest.raises(ValueError):calculate_shortline(ShortlineParams(axle_load=30.0),rules,catalog)
    with pytest.raises(ValueError):calculate_shortline(ShortlineParams(km=.05,yard_amvs=10),rules,catalog)


def test_parametros_salvos_sao_restaurados(model):
    rules,catalog=model
    params=ShortlineParams(km=42,sidings=3,siding_length=1800,yard_length=2500,yard_amvs=5,axle_load=20.0,
        drainage='Reforçada',fence='Muro',detection='Circuito de via',level_crossings=9,locomotives=4,wagons=60)
    result=calculate_shortline(params,rules,catalog)
    assert params_from_scenario(result['scenario'])==params


def test_excel_shortline_traz_premissas_e_grupo_de_patio(model):
    rules,catalog=model
    result=select_groups(calculate_shortline(ShortlineParams(level_crossings=2),rules,catalog),list(SHORTLINE_GROUPS))
    used={i['price_key']:catalog[i['price_key']] for i in result['items']}
    workbook=load_workbook(BytesIO(make_excel(result,used)))
    assert SIDING_GROUP in workbook.sheetnames
    premises=[row[1] for row in workbook['Premissas'].iter_rows(values_only=True)]
    assert 'Desvios de cruzamento' in premises and 'Passagens em nível' in premises
    memory=[row[11] for row in workbook['EAP'].iter_rows(min_row=5,values_only=True) if row[1]==SIDING_GROUP]
    assert memory and all('desvios e pátio' in text for text in memory)


def test_custo_por_km_de_via_acompanha_selecao_e_bdi(model):
    rules,catalog=model
    result=calculate_shortline(ShortlineParams(),rules,catalog,bdi=.263)
    track_km=result['scenario']['track_km']
    chosen=[g for g in SHORTLINE_GROUPS if default_selection(g)]
    with_bdi=select_groups(result,chosen,bdi_rate=.263)
    without_bdi=select_groups(result,chosen,bdi_rate=0.0)
    assert with_bdi['per_track_km']==pytest.approx(with_bdi['total']/track_km,abs=.01)
    assert without_bdi['per_track_km']==pytest.approx(without_bdi['direct']/track_km,abs=.01)
    assert with_bdi['per_track_km']<result['per_track_km']  # sinalização fora do recorte
    assert without_bdi['per_track_km']<with_bdi['per_track_km']
    summary=load_workbook(BytesIO(make_excel(with_bdi,catalog)))['Resumo']
    rows={row[0]:row[1] for row in summary.iter_rows(values_only=True) if row and row[0]}
    assert rows['R$/km de via (principal + desvios + pátio)']==with_bdi['per_track_km']
