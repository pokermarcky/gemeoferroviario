"""Integridade do Excel exportado: abas, valores de erro e fechamento dos totais."""
from io import BytesIO
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from railbudget.elevated import calculate_elevated
from railbudget.engine import Scenario, calculate, load_model
from railbudget.exporters import make_excel


ERRORS = ("#REF!", "#N/A", "#NAME?", "#VALUE!", "#DIV/0!", "#NUM!", "#NULL!")


@pytest.fixture(scope="module", params=["superficie", "elevado"])
def exported(request):
    rules, catalog = load_model()
    if request.param == "elevado":
        result = calculate_elevated(Scenario(configuration="Elevado", km=2.5, lines=2, amvs=2,
            rolling_stock=False, elevated_span=30.0), rules, catalog)
    else:
        result = calculate(Scenario(), rules, catalog)
    return result, make_excel(result, catalog)


def test_abas_esperadas_e_sem_valores_de_erro(exported):
    result, data = exported
    workbook = load_workbook(BytesIO(data))
    expected = {"Resumo", "EAP", "Premissas", "SIEC_SERVICOS", *(group[:31] for group in result["groups"])}
    if result.get("elevated"):
        expected.add("Elevado")
    assert expected <= set(workbook.sheetnames)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value.strip() in ERRORS), (sheet.title, cell.coordinate)


def test_eap_e_resumo_fecham_com_o_motor(exported):
    result, data = exported
    workbook = load_workbook(BytesIO(data), data_only=True)
    eap = workbook["EAP"]
    totals = [eap[f"I{row}"].value for row in range(5, eap.max_row + 1)]
    assert len(totals) == len(result["items"])
    assert round(sum(totals), 2) == result["direct"]
    summary = workbook["Resumo"]
    values = {summary[f"A{row}"].value: summary[f"B{row}"].value for row in range(1, summary.max_row + 1)}
    assert values["Direto"] == result["direct"]
    assert values["Total"] == result["total"]
    for group, direct in result["groups"].items():
        assert values[group] == direct


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice não instalado")
def test_libreoffice_abre_o_arquivo(exported, tmp_path):
    _, data = exported
    source = tmp_path / "orcamento.xlsx"
    source.write_bytes(data)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "csv", "--outdir", str(tmp_path), str(source)],
        check=True, capture_output=True, timeout=180)
    assert (tmp_path / "orcamento.csv").stat().st_size > 0
