from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from railbudget.comparison import compare_results, comparison_excel
from railbudget.engine import Scenario, calculate, load_model, select_groups
from railbudget.expressions import money


@pytest.fixture(scope='module')
def model():
    return load_model()


def test_simples_contra_dupla_por_grupo_e_total(model):
    rules,catalog=model
    a=select_groups(calculate(Scenario(km=2,lines=1),rules,catalog),['1 Via permanente','3 Drenagem'])
    b=select_groups(calculate(Scenario(km=2,lines=2),rules,catalog),['1 Via permanente','3 Drenagem'])
    before=deepcopy((a,b))
    rows={row['item']:row for row in compare_results(a,b)}
    assert (a,b)==before  # só leitura
    assert set(rows)>={'1 Via permanente','3 Drenagem','Total','Por km de corredor'}
    assert '4 Vedação' not in rows  # fora dos dois recortes
    via=rows['1 Via permanente']
    assert via['difference']==money(b['groups']['1 Via permanente']-a['groups']['1 Via permanente'])>0
    total=rows['Total']
    assert total['a']==a['total'] and total['b']==b['total']
    assert total['percent']==pytest.approx(100*(b['total']-a['total'])/a['total'])


def test_grupo_desmarcado_em_um_lado_vale_zero_e_variacao_indefinida(model):
    rules,catalog=model
    full=calculate(Scenario(),rules,catalog)
    a=select_groups(full,['1 Via permanente'])
    b=select_groups(full,['1 Via permanente','4 Vedação'])
    rows={row['item']:row for row in compare_results(a,b)}
    assert rows['4 Vedação']['a']==0 and rows['4 Vedação']['b']==full['groups']['4 Vedação']
    assert rows['4 Vedação']['percent'] is None
    assert rows['1 Via permanente']['difference']==0


def test_planilha_da_comparacao(model):
    rules,catalog=model
    a=calculate(Scenario(km=1),rules,catalog)
    b=calculate(Scenario(km=3),rules,catalog)
    rows=compare_results(a,b)
    sheet=load_workbook(BytesIO(comparison_excel(rows,'A: 1 km','B: 3 km')))['Comparação']
    values=list(sheet.iter_rows(values_only=True))
    assert values[1][1]=='A: 1 km' and values[4][0]=='Item'
    total=next(row for row in values if row and row[0]=='Total')
    assert total[1]==a['total'] and total[2]==b['total'] and total[3]==money(b['total']-a['total'])
