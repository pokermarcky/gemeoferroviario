from io import BytesIO
import pytest
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from railbudget.engine import *
from railbudget.exporters import export_all

def test_export_values_formulas_and_documents():
    rules,catalog=load_model()
    result=calculate(Scenario(configuration='Elevado',lines=2,amvs=3,km=1.3,drainage='Complexa'),rules,catalog)
    files=export_all(result,catalog)
    f=load_workbook(BytesIO(files['orcamento.xlsx']),data_only=False)
    v=load_workbook(BytesIO(files['orcamento.xlsx']),data_only=True)
    total_row=next(n for n in range(1,v['Resumo'].max_row+1) if v['Resumo'][f'A{n}'].value=='Total')
    assert v['Resumo'][f'B{total_row}'].value==result['total']
    assert len(f['Resumo']._charts)==1
    assert f['EAP']['H5'].value.startswith('=VLOOKUP(')
    assert f['EAP'].freeze_panes=='A5'
    assert f['Premissas']['A2'].value=='Parâmetro V5'
    assert 'fator de gestão' in f['EAP']['L5'].value
    for n,item in enumerate(result['items'],5):
        assert v['EAP'][f'I{n}'].value==item['total']
        assert v['EAP'][f'G{n}'].value==pytest.approx(item['quantity'])
    word=Document(BytesIO(files['relatorio.docx']))
    assert word.tables
    assert any('Memória de cálculo:' in paragraph.text for paragraph in word.paragraphs)
    for name in ['orcamento.pdf','relatorio.pdf']:
        doc=PdfReader(BytesIO(files[name]));assert len(doc.pages)>1
        text=''.join(page.extract_text() for page in doc.pages)
        assert all(item['code'] in text for item in result['items'])
        if name=='relatorio.pdf':assert 'Memória de cálculo:' in text
