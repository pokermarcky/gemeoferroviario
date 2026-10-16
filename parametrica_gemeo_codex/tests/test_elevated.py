from io import BytesIO

import pytest
from openpyxl import load_workbook

from railbudget.engine import Scenario, calculate, load_model
from railbudget.elevated import calculate_elevated, sensitivity
from railbudget.exporters import make_excel


@pytest.fixture(scope='module')
def model():
    return load_model()


def two_segments():
    return [
        {'Trecho':'Trecho baixo','Extensão (km)':.4,'Altura (m)':6,'Vão (m)':20,
            'Fundação':'Hélice contínua','Diâmetro (m)':.8,'Profundidade (m)':18,
            'Estacas/apoio':4,'Interferência':'Baixa'},
        {'Trecho':'Trecho alto','Extensão (km)':.6,'Altura (m)':14,'Vão (m)':30,
            'Fundação':'Estaca raiz','Diâmetro (m)':.41,'Profundidade (m)':25,
            'Estacas/apoio':8,'Interferência':'Alta'},
    ]


def test_modelo_avancado_preserva_referencia_e_segmenta_quantitativos(model):
    rules,catalog=model
    scenario=Scenario(configuration='Elevado',rolling_stock=False)
    base=calculate(scenario,rules,catalog)
    default=calculate_elevated(scenario,rules,catalog)
    assert default['total']==base['total']
    assert default['elevated']['technical_factors']['Tipologia do tabuleiro']['factor']==1

    segmented=calculate_elevated(scenario,rules,catalog,two_segments())
    summaries=segmented['elevated']['segment_summaries']
    assert [row['Apoios'] for row in summaries]==[21,20]
    assert segmented['elevated']['total_piles']==244
    assert segmented['elevated']['total_pile_length']==5512
    assert sum(item['label'].endswith('Mobilização fundações') for item in segmented['items'])==2
    assert any(item['label'].startswith('Trecho alto · Estacas') for item in segmented['items'])


def test_solucoes_e_provisoes_sao_declaradas_e_alteram_total(model):
    rules,catalog=model
    scenario=Scenario(configuration='Elevado',rolling_stock=False,
        deck_type='Viga-caixão pré-moldada segmental',pier_type='Pórtico',
        geotechnical_condition='Desfavorável',construction_method='Aduelas pré-moldadas',
        execution_constraint='Ambiente urbano controlado',special_span_length=100,
        transitions=2,transition_unit_cost=1500000,noise_barrier_length=500,
        noise_barrier_height=3,noise_barrier_unit_cost=2400)
    result=calculate_elevated(scenario,rules,catalog,two_segments())
    provisions=[item for item in result['items'] if item['id'].startswith('elevado:')]
    assert {item['label'] for item in provisions}=={
        'Adicional para vãos especiais','Encontros e transições','Barreira acústica'}
    assert result['custom_prices']
    assert result['elevated']['indicators']['Trechos homogêneos']==2
    assert result['elevated']['technical_factors']['Tipologia dos pilares']['factor']==1.28
    assert len(sensitivity(result,scenario,rules,catalog,two_segments()))==4


def test_excel_avancado_documenta_trechos_fatores_indicadores_e_provisoes(model):
    rules,catalog=model
    scenario=Scenario(configuration='Elevado',rolling_stock=False,transitions=1,
        noise_barrier_length=100,deck_type='Viga U pré-moldada')
    result=calculate_elevated(scenario,rules,catalog,two_segments())
    workbook=load_workbook(BytesIO(make_excel(result,catalog)),data_only=True)
    elevated=workbook['Elevado']
    categories={elevated[f'A{row}'].value for row in range(5,elevated.max_row+1)}
    assert {'Trecho homogêneo','Coeficiente técnico','Indicador gerencial','Custo estrutural'}.issubset(categories)
    eap_text=' '.join(str(cell.value or '') for row in workbook['EAP'].iter_rows() for cell in row)
    assert 'Encontros e transições' in eap_text
    assert 'Barreira acústica' in eap_text
    assert all(sheet.freeze_panes is None for sheet in workbook.worksheets)


def test_soma_dos_trechos_deve_reconciliar_com_corredor(model):
    rules,catalog=model
    scenario=Scenario(configuration='Elevado',rolling_stock=False)
    invalid=two_segments();invalid[1]['Extensão (km)']=.5
    with pytest.raises(ValueError,match='soma dos trechos'):
        calculate_elevated(scenario,rules,catalog,invalid)
