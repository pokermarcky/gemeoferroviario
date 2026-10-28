"""Desenho do elevado nos arquivos baixados (Excel, Word e PDFs)."""
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from railbudget import elevated_drawing
from railbudget.drawing_image import elevated_drawing_png, svg_to_png
from railbudget.engine import Scenario, calculate, load_model
from railbudget.elevated import DECK_FACTORS, PIER_FACTORS
from railbudget.exporters import export_all


@pytest.fixture(scope='module')
def model():
    return load_model()


def elevated_result(model,**changes):
    rules,catalog=model
    return calculate(Scenario(configuration='Elevado',lines=2,amvs=3,km=1.3,**changes),rules,catalog)


def colors_of(png):
    image=Image.open(BytesIO(png)).convert('RGB')
    return {color for _,color in image.getcolors(image.size[0]*image.size[1])}


@pytest.mark.parametrize('deck',list(DECK_FACTORS))
@pytest.mark.parametrize('pier',list(PIER_FACTORS))
def test_desenho_vira_png_em_todas_as_tipologias(model,deck,pier):
    png=elevated_drawing_png(elevated_result(model,deck_type=deck,pier_type=pier))
    image=Image.open(BytesIO(png))
    assert image.format=='PNG' and image.size[0]==2020
    assert (0xb8,0x32,0x2a) in colors_of(png)  # trilhos em destaque


def test_png_usa_as_cores_do_desenho_da_tela(model,monkeypatch):
    """Se a cor do desenho mudar (ex.: tema mais suave), o arquivo baixado acompanha."""
    monkeypatch.setattr(elevated_drawing,'RAIL','#123456')
    monkeypatch.setattr(elevated_drawing,'VIA_LEGEND',[('#123456','Trilhos')]+elevated_drawing.VIA_LEGEND[1:])
    png=elevated_drawing_png(elevated_result(model))
    assert (0x12,0x34,0x56) in colors_of(png)
    assert (0xb8,0x32,0x2a) not in colors_of(png)


def test_sem_elevado_nao_ha_desenho(model):
    rules,catalog=model
    result=calculate(Scenario(lines=1,km=2),rules,catalog)
    assert elevated_drawing_png(result) is None
    files=export_all(result,catalog)
    assert 'Desenho do elevado' not in load_workbook(BytesIO(files['orcamento.xlsx'])).sheetnames
    assert not Document(BytesIO(files['relatorio.docx'])).inline_shapes
    for name in ('orcamento.pdf','relatorio.pdf'):
        assert 'Desenho do elevado' not in ''.join(page.extract_text() for page in PdfReader(BytesIO(files[name])).pages)


def test_downloads_do_elevado_incluem_desenho_e_via_permanente(model):
    rules,catalog=model
    result=elevated_result(model)
    files=export_all(result,catalog)
    book=load_workbook(BytesIO(files['orcamento.xlsx']))
    assert book.sheetnames[-1]=='Desenho do elevado' and len(book['Desenho do elevado']._images)==1
    track=[row for row in book['Elevado'].iter_rows(values_only=True) if row[0]=='Via permanente no elevado']
    assert track and track[-1][1].startswith('Total da via permanente')
    word=Document(BytesIO(files['relatorio.docx']))
    assert len(word.inline_shapes)==1
    headings=[p.text for p in word.paragraphs if p.style.name.startswith('Heading')]
    assert 'Desenho do elevado' in headings and 'Via permanente no elevado' in headings
    for name in ('orcamento.pdf','relatorio.pdf'):
        reader=PdfReader(BytesIO(files[name]))
        text=''.join(page.extract_text() for page in reader.pages)
        assert 'Desenho do elevado' in text and 'Via permanente no elevado' in text
        assert sum(len(page.images) for page in reader.pages)==1


def test_desenho_nao_muda_valores(model):
    rules,catalog=model
    result=elevated_result(model)
    total=result['total'];direct=result['direct']
    export_all(result,catalog)
    assert (result['total'],result['direct'])==(total,direct)


def test_svg_desconhecido_minimo():
    png=svg_to_png('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 5"><rect x="0" y="0" width="10" height="5" fill="#00ff00"/></svg>',1)
    assert Image.open(BytesIO(png)).size==(10,5)
