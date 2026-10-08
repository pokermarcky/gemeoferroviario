from io import BytesIO

from openpyxl import load_workbook

from railbudget.vlt import (BASE_TOTAL, calculate_vlt, make_vlt_excel)


def test_vlt_reproduz_total_do_edital_e_extensoes_hibridas():
    result=calculate_vlt()
    assert result['total']==BASE_TOTAL
    assert result['scenario']['surface_km']==.9
    assert result['scenario']['elevated_km']==1.7
    assert result['scenario']['km']==2.6
    assert result['scenario']['line_km']==5.2
    assert result['per_km']==round(BASE_TOTAL/2.6,2)
    assert sum(result['segments'].values())==BASE_TOTAL
    assert len(result['items'])==14


def test_vlt_parametriza_superficie_e_elevado_separadamente():
    base=calculate_vlt();surface=calculate_vlt(1.8,1.7);elevated=calculate_vlt(.9,3.4)
    assert surface['segments']['Trecho em superfície']==2*base['segments']['Trecho em superfície']
    assert surface['segments']['Trecho elevado']==base['segments']['Trecho elevado']
    assert elevated['segments']['Trecho elevado']==2*base['segments']['Trecho elevado']
    assert elevated['segments']['Trecho em superfície']==base['segments']['Trecho em superfície']
    assert surface['segments']['Custos fixos']==base['segments']['Custos fixos']


def test_excel_vlt_e_claro_e_abre_sem_reparo():
    result=calculate_vlt();workbook=load_workbook(BytesIO(make_vlt_excel(result)),data_only=True)
    assert workbook.sheetnames==['Resumo','Grupos do edital','Premissas']
    assert all(sheet.freeze_panes is None for sheet in workbook.worksheets)
    assert workbook['Resumo']['B2'].value=='VLT Aeroporto - Castelão'
    assert workbook['Resumo']['B9'].value==BASE_TOTAL
    assert workbook['Premissas']['B12'].value==BASE_TOTAL
    assert workbook['Grupos do edital'].max_row==15
