from copy import deepcopy
from io import BytesIO
import pytest
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from railbudget.engine import Scenario,load_model,calculate,select_groups
from railbudget.expressions import money
from railbudget.exporters import export_all,currency


@pytest.fixture
def model():
    rules,catalog=load_model()
    return calculate(Scenario(km=2,fence='Muro'),rules,catalog),catalog


def test_individual_group_retains_quantities_and_applies_bdi_once(model):
    full,_=model;before=deepcopy(full)
    chosen=select_groups(full,['4 Vedação'])
    assert chosen['direct']==full['groups']['4 Vedação']
    assert chosen['total']==money(chosen['direct']+money(chosen['direct']*full['scenario']['bdi']))
    assert chosen['per_km']==money(chosen['total']/2)
    assert chosen['items']==[x for x in full['items'] if x['group']=='4 Vedação']
    assert full==before
    assert next(x for x in chosen['scope_summary'] if x['group']=='4 Vedação')['share']==100


def test_multiple_none_and_all_groups(model):
    full,_=model
    some=select_groups(full,['3 Drenagem','6 Banco de dutos'])
    assert some['direct']==money(full['groups']['3 Drenagem']+full['groups']['6 Banco de dutos'])
    assert sum(x['share'] for x in some['scope_summary'])==pytest.approx(100)
    assert select_groups(full,full['groups'])['total']==full['total']
    empty=select_groups(full,[])
    assert empty['total']==empty['per_km']==0
    assert not empty['items'] and not empty['counts']
    assert all(x['share']==0 for x in empty['scope_summary'])
    with pytest.raises(ValueError):select_groups(full,['Grupo inexistente'])


def test_custom_bdi_reconciles_groups_and_export_context(model):
    full,_=model
    chosen=select_groups(full,full['groups'],bdi_rate=0.25)
    assert chosen['bdi_amount']==money(full['direct']*.25)
    assert chosen['scenario']['bdi']==chosen['context']['bdi']==.25
    assert money(sum(g['total'] for g in chosen['scope_summary']))==chosen['total']
    assert full['scenario']['bdi']!=.25
    off=select_groups(full,full['groups'],bdi_rate=0)
    assert off['total']==off['direct'] and off['bdi_amount']==0
    for rate in [-.01,1.01,float('nan')]:
        with pytest.raises(ValueError):select_groups(full,full['groups'],bdi_rate=rate)


def test_segregated_excel_and_selected_documents(model):
    full,catalog=model
    chosen=select_groups(full,['4 Vedação','6 Banco de dutos'])
    files=export_all(chosen,catalog)
    values=load_workbook(BytesIO(files['orcamento.xlsx']),data_only=True)
    formulas=load_workbook(BytesIO(files['orcamento.xlsx']),data_only=False)
    assert values['Resumo']['B13'].value==chosen['total']
    assert sum(values[g]['C6'].value for g in full['groups'])==chosen['direct']
    for g in full['groups']:
        sh=values[g]
        assert sh['B2'].value==g and sh['A1'].value is None
        assert sh['C6'].value==chosen['groups'][g]
        assert sh['C6'].font.name=='Aptos' and sh['C6'].font.sz==12
        rows=[(n,x) for n,x in enumerate(chosen['items'],5) if x['group']==g]
        for dest,(source,item) in enumerate(rows,13):
            assert formulas[g][f'I{dest}'].value==f"='EAP'!I{source}"
            assert sh[f'I{dest}'].value==item['total']
        if not rows:assert sh['B4'].value=='Excluído do total selecionado'
    word=Document(BytesIO(files['relatorio.docx']))
    assert currency(chosen['total']) in '\n'.join(c.text for t in word.tables for row in t.rows for c in row.cells)
    for name in ['orcamento.pdf','relatorio.pdf']:
        text=''.join(p.extract_text() for p in PdfReader(BytesIO(files[name])).pages)
        assert currency(chosen['total']) in text
        assert 'Recorte financeiro' in text
